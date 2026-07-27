import argparse
import csv
import re
import sys
from pathlib import Path

import anthropic

try:
    import docx  # python-docx
except ImportError:
    docx = None

from count_letter_categories import read_text

DEFAULT_MODEL = "claude-opus-4-8"
SAFE_CHARS_RE = re.compile(r"[^A-Za-z0-9_.-]+")


def build_prompt(template: str, candidate_name: str, resume_text: str) -> str:
    return template.format(candidate_name=candidate_name, resume_text=resume_text)


def generate_letter(client, model, effort, prompt) -> str:
    response = client.messages.create(
        model=model,
        max_tokens=4096,
        thinking={"type": "adaptive"},
        output_config={"effort": effort},
        messages=[{"role": "user", "content": prompt}],
    )
    return next(b.text for b in response.content if b.type == "text")


def write_docx(text: str, path: Path) -> None:
    doc = docx.Document()
    for para in text.split("\n\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    doc.save(str(path))


def safe_filename(text: str) -> str:
    return SAFE_CHARS_RE.sub("_", text.strip())


def main():
    parser = argparse.ArgumentParser(
        description="Generate one recommendation letter per name (grouped by category) "
        "from a single resume, using the Claude API. Useful for testing whether "
        "letter language (per dictionary_v3.csv categories) varies by candidate name."
    )
    parser.add_argument("resume", help="Path to the single resume file (.txt or .docx) used for every letter")
    parser.add_argument(
        "--names-csv",
        required=True,
        help="CSV with columns: category, name -- one row per candidate name to generate a letter for",
    )
    parser.add_argument(
        "--prompt-template",
        default="prompt_template.txt",
        help="Path to a text file with the prompt. Supports {candidate_name} and {resume_text} placeholders.",
    )
    parser.add_argument(
        "--output-dir",
        default="generated_letters",
        help="Directory to write generated .docx letters and manifest.csv into",
    )
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument(
        "--effort",
        default="medium",
        choices=["low", "medium", "high", "xhigh", "max"],
        help="Claude effort level (default: medium)",
    )
    args = parser.parse_args()

    if docx is None:
        sys.exit("python-docx is not installed. Run: pip install python-docx")

    template = Path(args.prompt_template).read_text(encoding="utf-8")
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    resume_text = read_text(args.resume)

    client = anthropic.Anthropic()

    with open(args.names_csv, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    manifest = []
    failures = []
    for i, row in enumerate(rows, 1):
        category = row.get("category", "").strip()
        name = row.get("name", "").strip()
        print(f"[{i}/{len(rows)}] {name} ({category})")

        try:
            prompt = build_prompt(template, name, resume_text)
            letter_text = generate_letter(client, args.model, args.effort, prompt)
        except Exception as e:
            print(f"  FAILED ({type(e).__name__}): {e}", file=sys.stderr)
            failures.append(name or "?")
            continue

        filename = f"{safe_filename(category)}__{safe_filename(name)}.docx"
        out_path = output_dir / filename
        write_docx(letter_text, out_path)
        manifest.append({**row, "filename": filename})
        print(f"  wrote {out_path}")

    manifest_path = output_dir / "manifest.csv"
    manifest_fields = list(rows[0].keys()) + ["filename"] if rows else ["category", "name", "filename"]
    with open(manifest_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=manifest_fields)
        writer.writeheader()
        writer.writerows(manifest)

    print(f"\nDone. {len(manifest)}/{len(rows)} letters written to {output_dir}/ (manifest.csv written)")
    if failures:
        print(f"Failed: {', '.join(failures)}", file=sys.stderr)


if __name__ == "__main__":
    main()
